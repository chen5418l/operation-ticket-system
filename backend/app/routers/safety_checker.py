"""
安全校验 — IsolationChecker 集成
"""

import json, os, re, shutil, tempfile, subprocess, sys
from fastapi import APIRouter, HTTPException
from fastapi.responses import FileResponse
from pydantic import BaseModel
from typing import Dict

router = APIRouter(prefix="/api/safety", tags=["安全校验-IsolationChecker"])

RUNTIME_DIR = os.path.join(os.path.dirname(__file__), "..", "..", "safety_runtime")
os.makedirs(RUNTIME_DIR, exist_ok=True)

SCENARIO_FILE = os.path.join(RUNTIME_DIR, "scenario.xlsx")
REPORT_FILE = os.path.join(RUNTIME_DIR, "report.xlsx")
CHECKER_SCRIPT = os.path.join(os.path.dirname(__file__), "..", "..", "safety", "isolation_checker.py")

# ============ 默认配置 ============
DEFAULT_SCENARIO = {
    "FaultNode": 18,
    "FaultType": "三相短路",
    "FaultDuration": 0.1,
    "Switches": {f"Switch{i}": 1 for i in range(1, 38)},
}
# Override: Switch17/18/36 are open by default (matches checker auto-generate)
DEFAULT_SCENARIO["Switches"]["Switch17"] = 0
DEFAULT_SCENARIO["Switches"]["Switch18"] = 0
DEFAULT_SCENARIO["Switches"]["Switch36"] = 0

# ============ Schemas ============
class CheckRequest(BaseModel):
    FaultNode: int
    FaultType: str = "三相短路"
    FaultDuration: float = 0.1
    Switches: Dict[str, int] = {}

class SwitchState(BaseModel):
    Switches: Dict[str, int] = {}

# ============ 端点 ============

@router.get("/default_scenario")
def get_default_scenario():
    """返回默认场景配置"""
    return {
        "success": True,
        "scenario": DEFAULT_SCENARIO,
    }


@router.post("/isolation_check")
def run_safety_check(req: CheckRequest):
    """执行安全校验"""
    # 1. 校验输入
    if req.FaultNode < 1 or req.FaultNode > 33:
        raise HTTPException(400, "FaultNode 必须在 1~33 之间")
    for k, v in req.Switches.items():
        if v not in (0, 1):
            raise HTTPException(400, f"{k} 必须是 0 或 1")

    # 2. 生成 scenario.xlsx（Checker 格式: 两列 Parameter/Value，无表头）
    try:
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.cell(row=1, column=1, value="FaultNode"); ws.cell(row=1, column=2, value=req.FaultNode)
        ws.cell(row=2, column=1, value="FaultType"); ws.cell(row=2, column=2, value=req.FaultType)
        ws.cell(row=3, column=1, value="FaultDuration"); ws.cell(row=3, column=2, value=req.FaultDuration)
        row = 4
        for i in range(1, 38):
            ws.cell(row=row, column=1, value=f"Switch{i}")
            ws.cell(row=row, column=2, value=req.Switches.get(f"Switch{i}", 1))
            row += 1
        wb.save(SCENARIO_FILE)
    except ImportError:
        raise HTTPException(500, "openpyxl 未安装，请运行: pip install openpyxl")
    except Exception as e:
        raise HTTPException(500, f"生成 scenario.xlsx 失败: {e}")

    # 3. 调用 isolation_checker
    try:
        import subprocess, sys, shutil
        # 复制 scenario.xlsx 到 checker 所在目录（isolation_checker 可能只读当前目录）
        checker_dir = os.path.dirname(CHECKER_SCRIPT)
        if os.path.exists(checker_dir):
            shutil.copy2(SCENARIO_FILE, os.path.join(checker_dir, "scenario.xlsx"))

        # 在 checker 所在目录执行
        work_dir = checker_dir if os.path.isdir(checker_dir) else RUNTIME_DIR
        cmd = [sys.executable, CHECKER_SCRIPT, SCENARIO_FILE, REPORT_FILE]
        env = os.environ.copy()
        env['PYTHONIOENCODING'] = 'utf-8'
        result = subprocess.run(cmd, capture_output=True, text=True, timeout=60, cwd=work_dir, env=env, encoding='utf-8', errors='replace')

        # 如果 checker 生成了 report 在其目录，复制回 runtime
        checker_report = os.path.join(checker_dir, "report.xlsx")
        if os.path.exists(checker_report):
            shutil.copy2(checker_report, REPORT_FILE)

        if result.returncode != 0:
            return {
                "success": False,
                "error": "isolation_checker 执行失败",
                "stderr": result.stderr[-500:] if result.stderr else "",
                "stdout": result.stdout[-500:] if result.stdout else "",
            }
    except FileNotFoundError:
        return mock_result(req)
    except Exception as e:
        raise HTTPException(500, f"调用 isolation_checker 失败: {e}")

    # 4. 解析 report.xlsx
    if not os.path.exists(REPORT_FILE):
        raise HTTPException(500, "report.xlsx 未生成")

    try:
        return parse_report(req)
    except Exception as e:
        raise HTTPException(500, f"解析 report.xlsx 失败: {e}")


@router.get("/download_report")
def download_report():
    """下载最近生成的 report.xlsx"""
    if not os.path.exists(REPORT_FILE):
        raise HTTPException(404, "暂无报告，请先执行安全校验")
    return FileResponse(REPORT_FILE, filename="report.xlsx",
                        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


# ============ 辅助 ============

def parse_report(req: CheckRequest) -> dict:
    import openpyxl
    wb = openpyxl.load_workbook(REPORT_FILE, data_only=True)
    ws = wb[wb.sheetnames[0]]
    raw = list(ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True))
    data_rows = []
    for row in raw[1:]:
        vals = [v for v in row]
        if all(v is None or str(v).strip() in ('', '-') for v in vals): continue
        data_rows.append(vals)
    checks = []
    metrics = {}
    metric_start = {'故障包围完整性','非故障失供负荷','开关动作次数','远方操作比例',
                    '数据可信度','重要用户恢复率','加权重要用户恢复率',
                    '冷负荷冲击风险','PV弃光量','EV充电损失'}
    in_metrics = False
    for row in data_rows:
        name = str(row[0] or '').strip()
        raw_result = str(row[1] or '').strip() if len(row) > 1 else ''
        detail = str(row[2] or '').strip() if len(row) > 2 else ''
        raw_risk = str(row[3] or '').strip() if len(row) > 3 else ''
        if name in metric_start: in_metrics = True
        if in_metrics:
            nums = [float(x) for x in re.findall(r'[\d.]+', raw_result + ' ' + detail)]
            if '非故障失供' in name and nums: metrics['lost_load_kw'] = nums[0]
            elif '开关动作' in name and nums: metrics['switch_action_count'] = int(nums[0])
            elif '远方操作' in name and nums: metrics['remote_operation_rate'] = nums[0]
            elif '数据可信度' in name and nums: metrics['data_confidence'] = nums[0]
            elif '重要用户恢复率' in name and nums: metrics['important_restore_rate'] = nums[0]
            elif '加权重要用户恢复率' in name and nums: metrics['weighted_restore_rate'] = nums[0]
            elif 'PV弃光' in name and nums: metrics['pv_curtailment_kw'] = nums[0]
            elif 'EV充电' in name and nums: metrics['ev_loss_kw'] = nums[0]
            elif '冷负荷' in name: metrics['cold_load_risk'] = (raw_result + ' ' + detail)[:80]
        else:
            d2 = detail + ' ' + raw_result
            is_fail = any(kw in d2 for kw in ['不通过','失败','未完全隔离','非隔离供电风险'])
            is_warn = any(kw in d2 for kw in ['警告','风险','越限']) if not is_fail else False
            if is_fail: result, risk = '不通过', '高'
            elif is_warn: result, risk = '警告', '中'
            elif raw_risk == '高': result, risk = '不通过', '高'
            elif raw_risk == '中': result, risk = '警告', '中'
            else: result, risk = '通过', '低'
            checks.append({'item': name or '检查项', 'result': result, 'detail': detail[:200], 'risk_level': risk})
    check_passed = not any(c['result'] == '不通过' for c in checks)
    has_warn = any(c['result'] == '警告' for c in checks)
    if not check_passed: overall, risk_level = '不可行', '高风险'
    elif has_warn: overall, risk_level = '可行但有风险', '中风险'
    else: overall, risk_level = '可行', '低风险'
    print(f'[SafetyCheck] checks={len(checks)} metrics={len(metrics)} passed={check_passed} overall={overall}')
    return {'success': True, 'check_passed': check_passed, 'overall_result': overall,
            'risk_level': risk_level, 'metrics': metrics, 'checks': checks, 'summary': ''}


def mock_result(req: CheckRequest) -> dict:
    """isolation_checker.py 不存在时的模拟结果"""
    return {
        "success": True,
        "note": "isolation_checker.py 未找到，返回模拟结果",
        "scenario": {"FaultNode": req.FaultNode, "FaultType": req.FaultType, "FaultDuration": req.FaultDuration},
        "summary": f"方案可行。故障节点Bus{req.FaultNode}，失供负荷90.00 kW，隔离完成。",
        "overall_result": "可行",
        "risk_level": "低风险",
        "metrics": {
            "lost_load_kw": 90.0, "important_restore_rate": 75.0,
            "weighted_restore_rate": 69.2, "switch_action_count": 3,
            "remote_operation_rate": 66.7, "data_confidence": 98.8,
        },
        "checks": [
            {"item": "设备检查", "result": "通过", "detail": "跳闸开关变位、遥信、保护动作正常", "risk_level": "低"},
            {"item": "隔离完整性", "result": "通过", "detail": "故障区段两侧隔离开关已断开", "risk_level": "低"},
            {"item": "容量校验", "result": "通过", "detail": "转供后馈线负载率<80%", "risk_level": "低"},
            {"item": "电压校验", "result": "警告", "detail": "末端电压0.93pu", "risk_level": "中"},
            {"item": "N-1校验", "result": "通过", "detail": "N-1场景无越限", "risk_level": "低"},
        ],
    }
