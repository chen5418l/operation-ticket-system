"""操作票 API"""
import json, io, datetime
from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from pydantic import BaseModel
from app.core.database import get_db
from app.models.ticket import OperationTicket
from app.services.ticket_service import generate_ticket

router = APIRouter(prefix="/api", tags=["操作票"])


class TicketRequest(BaseModel):
    planId: int
    operator: str = "操作人"
    guardian: str = "监护人"


@router.post("/ticket/generate")
def ticket_generate(req: TicketRequest, db: Session = Depends(get_db)):
    result = generate_ticket(db, req.planId, req.operator, req.guardian)
    return {"code": 200, "message": "success", "data": result}


@router.get("/tickets")
def list_tickets(db: Session = Depends(get_db)):
    tickets = db.query(OperationTicket).order_by(OperationTicket.created_at.desc()).all()
    return {
        "code": 200,
        "data": [{"id": t.id, "ticketNo": t.ticket_no, "taskName": t.task_name,
                   "operator": t.operator, "guardian": t.guardian,
                   "status": t.status, "checkResult": t.check_result,
                   "createdAt": str(t.created_at)} for t in tickets],
    }


@router.get("/ticket/{ticket_id}")
def get_ticket_detail(ticket_id: str, db: Session = Depends(get_db)):
    # 支持数字 ID 或票号字符串
    t = None
    if ticket_id.isdigit():
        t = db.query(OperationTicket).filter(OperationTicket.id == int(ticket_id)).first()
    if not t:
        t = db.query(OperationTicket).filter(OperationTicket.ticket_no == ticket_id).first()
    if not t:
        return {"code": 404, "message": "操作票不存在", "data": None}
    return {
        "code": 200,
        "data": {
            "id": t.id, "ticketNo": t.ticket_no, "taskName": t.task_name,
            "steps": json.loads(t.operation_steps) if isinstance(t.operation_steps, str) else [],
            "safetyNotes": json.loads(t.safety_notes) if isinstance(t.safety_notes, str) else [],
            "operator": t.operator, "guardian": t.guardian,
            "checkResult": t.check_result, "status": t.status,
            "createdAt": str(t.created_at),
        },
    }


@router.get("/ticket/{ticket_id}/export")
def export_ticket(ticket_id: str, db: Session = Depends(get_db)):
    """导出操作票为 Excel 文件"""
    t = None
    if ticket_id.isdigit():
        t = db.query(OperationTicket).filter(OperationTicket.id == int(ticket_id)).first()
    if not t:
        t = db.query(OperationTicket).filter(OperationTicket.ticket_no == ticket_id).first()
    if not t:
        return {"code": 404, "message": "操作票不存在"}

    try:
        import openpyxl
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, Border, Side, PatternFill

        wb = Workbook()
        thin = Border(left=Side('thin'), right=Side('thin'), top=Side('thin'), bottom=Side('thin'))
        hf = Font(bold=True, size=11)
        hfill = PatternFill(start_color="E8F0FE", end_color="E8F0FE", fill_type="solid")

        # Sheet 1: 基本信息
        ws1 = wb.active
        ws1.title = "操作票基本信息"
        info = [
            ("操作票编号", t.ticket_no), ("操作任务", t.task_name),
            ("操作人", t.operator), ("监护人", t.guardian),
            ("校验结论", t.check_result), ("状态", t.status),
            ("生成时间", str(t.created_at)),
        ]
        ws1.column_dimensions['A'].width = 18; ws1.column_dimensions['B'].width = 55
        for c, h in enumerate(["字段","内容"],1):
            cell = ws1.cell(row=1, column=c, value=h); cell.font = hf; cell.fill = hfill; cell.border = thin
        for r, (k, v) in enumerate(info, 2):
            ws1.cell(row=r, column=1, value=k).font = Font(bold=True)
            ws1.cell(row=r, column=1).border = thin
            ws1.cell(row=r, column=2, value=v).border = thin

        # Sheet 2: 操作步骤
        ws2 = wb.create_sheet("操作步骤")
        headers = ["序号","操作内容","安全提示"]
        widths = [6, 50, 40]
        for c, (h, w) in enumerate(zip(headers, widths), 1):
            cell = ws2.cell(row=1, column=c, value=h); cell.font = hf; cell.fill = hfill; cell.border = thin
            ws2.column_dimensions[cell.column_letter].width = w
        steps = json.loads(t.operation_steps) if isinstance(t.operation_steps, str) else []
        notes = json.loads(t.safety_notes) if isinstance(t.safety_notes, str) else []
        for i, s in enumerate(steps, 1):
            ws2.cell(row=i+1, column=1, value=i).border = thin
            ws2.cell(row=i+1, column=2, value=s).border = thin
            note = notes[i-1] if i-1 < len(notes) else ""
            ws2.cell(row=i+1, column=3, value=note).border = thin

        # Sheet 3: 安全提示
        ws3 = wb.create_sheet("安全注意事项")
        ws3.column_dimensions['A'].width = 55
        ws3.cell(row=1, column=1, value="安全注意事项").font = hf
        ws3.cell(row=1, column=1).fill = hfill; ws3.cell(row=1, column=1).border = thin
        for i, n in enumerate(notes, 2):
            ws3.cell(row=i, column=1, value=n).border = thin

        buf = io.BytesIO()
        wb.save(buf); buf.seek(0)
        now = datetime.datetime.now()
        from urllib.parse import quote
        filename = f"操作票_{t.ticket_no}_{now.strftime('%Y%m%d%H%M%S')}.xlsx"
        encoded_filename = quote(filename)
        return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                 headers={"Content-Disposition": f"attachment; filename*=UTF-8''{encoded_filename}"})
    except ImportError as e:
        return {"code": 500, "message": f"openpyxl 未安装: {e}"}
    except Exception as e:
        import traceback
        return {"code": 500, "message": f"导出失败: {e}", "traceback": traceback.format_exc()}
